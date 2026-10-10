# -*- coding: utf-8 -*-
"""生成《商业计划书设计纲要与撰写样板》（全文档宋体）。

复用 build_brief.py 中的排版原语（字体强制宋体、表格 keep-together、
字符画段中不分页等），只替换内容与封面。

  商业计划书设计纲要与样板.docx   ← 主交付物（宋体）
  商业计划书设计纲要与样板.md     ← 可编辑纯文本源
"""

import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import build_brief as BB  # noqa: E402
from bp_content import BLOCKS  # noqa: E402

ROOT = os.path.dirname(_HERE)
DOCX_OUT = os.path.join(ROOT, '商业计划书设计纲要与样板.docx')
MD_OUT = os.path.join(ROOT, '商业计划书设计纲要与样板.md')


def _cover(doc):
    for _ in range(3):
        BB.add_paragraph(doc, '', indent=False, space_after=0)
    BB.add_paragraph(doc, '挑 战 杯 参 赛 项 目', size=14, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                     space_after=24)
    BB.add_paragraph(doc, '商 业 计 划 书', size=30, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                     space_after=6)
    BB.add_paragraph(doc, '设 计 纲 要 与 撰 写 样 板', size=18, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                     space_after=16)
    BB.add_paragraph(doc, '基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统',
                     size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
                     indent=False, space_after=6)
    BB.add_paragraph(doc, '安步 · 让跌倒风险在迈出下一步之前被看见',
                     size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                     indent=False, space_after=18)
    BB.add_rule(doc)
    BB.add_paragraph(doc, '本 文 档 用 途', size=12, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                     space_after=8)
    BB.add_paragraph(
        doc,
        '本文件分两部分。上半部分是**设计纲要**：把商业计划书拆成逐章的'
        '写作任务，', size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER,
        indent=False, space_after=4, line=1.5)
    BB.add_paragraph(
        doc,
        '说明每一章要立住哪一句结论、写多长、用什么图、素材取自哪里、'
        '以及评委最容易在哪里扣分。', size=11.5,
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=4, line=1.5)
    BB.add_paragraph(
        doc,
        '下半部分是**简略样板**：一份可以直接在 Word 里扩写的完整草稿，'
        '每章都已写好骨架与语气，', size=11.5,
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=4, line=1.5)
    BB.add_paragraph(
        doc,
        '只需把【待填】【待实测】处补上你们自己的数据即可成稿。',
        size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
        space_after=24, line=1.5)
    for label in ('项目名称：基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统',
                  '建议品牌名：安步（待团队确认）',
                  '申报单位：【待填：学校 / 学院】',
                  '团队名称：【待填：团队名称】',
                  '项目负责人：【待填：姓名】',
                  '指导教师：【待填：姓名 / 职称】',
                  '编制日期：2026 年 10 月'):
        BB.add_paragraph(doc, label, size=12,
                         align=WD_ALIGN_PARAGRAPH.CENTER,
                         indent=False, space_after=6)


def _toc(doc):
    BB.page_section(doc)
    BB.add_paragraph(doc, '目  录', size=16, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                     space_after=12)
    BB.add_paragraph(doc, '提示：目录已自动生成。若后续修改正文导致页码变化，'
                          '请右键目录区域选择"更新域"→"更新整个目录"刷新页码。',
                     size=10.5, indent=False, space_after=10,
                     color=RGBColor(0x80, 0x80, 0x80))
    BB.add_toc(doc)


def render_docx():
    # 重置 build_brief 中的渲染状态（该模块的全局计数器是复用的）
    BB._PENDING['pb'] = False
    BB._PENDING['gap'] = False
    BB._H1_SEEN['v'] = False
    del BB._TABLE_LOG[:]

    doc = Document()
    BB.force_song_everywhere(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    for blk in BLOCKS:
        kind = blk[0]

        if kind == 'cover':
            _cover(doc)
        elif kind == 'toc':
            _toc(doc)
        elif kind == 'h1':
            if not BB._PENDING['pb'] and not BB._seen_h1_yet(doc):
                BB.page_section(doc)
                BB.add_page_number_footer(doc.sections[-1])
                BB.set_page_start(doc.sections[-1], 1)
            BB.add_heading(doc, blk[1], level=1)
        elif kind == 'h2':
            BB.add_heading(doc, blk[1], level=2)
        elif kind == 'h3':
            BB.add_heading(doc, blk[1], level=3)
        elif kind == 'p':
            BB.add_paragraph(doc, blk[1], bold=blk[2], indent=blk[3],
                             size=blk[4])
        elif kind == 'pn':
            BB.add_paragraph(doc, blk[1], bold=blk[2], indent=False,
                             size=blk[3])
        elif kind == 'b':
            BB.add_bullet(doc, blk[1])
        elif kind == 'table':
            _, headers, rows, widths, caption, size, keep = blk
            BB.add_table(doc, headers, rows, widths, size, caption, keep)
        elif kind == 'code':
            BB.add_code(doc, blk[1], size=blk[2])
        elif kind == 'center':
            BB.add_paragraph(doc, blk[1], size=blk[2], bold=blk[3],
                             align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                             space_before=6, space_after=8)
        elif kind == 'qa':
            _, tag, q, ans = blk
            qp = doc.add_paragraph()
            qf = qp.paragraph_format
            qf.line_spacing = 1.3
            qf.space_before = Pt(8)
            qf.space_after = Pt(3)
            qf.first_line_indent = Pt(0)
            qf.keep_with_next = True
            BB.add_runs(qp, tag + '　' + q, size=12, bold=True)
            ap = doc.add_paragraph()
            af = ap.paragraph_format
            af.line_spacing = 1.3
            af.space_after = Pt(4)
            af.left_indent = Pt(24)
            af.first_line_indent = Pt(0)
            BB.add_runs(ap, '答：' + ans, size=11.5)
        elif kind == 'pb':
            BB._PENDING['pb'] = True
        elif kind == 'end':
            BB.add_rule(doc)
            BB.add_paragraph(doc, '（文档结束）', size=11,
                             align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                             space_before=10)

    BB.enable_update_fields(doc)
    doc.save(DOCX_OUT)
    kept = [c for n, c, r in BB._TABLE_LOG if n <= BB.KEEP_TOGETHER_LINES]
    span = [c for n, c, r in BB._TABLE_LOG if n > BB.KEEP_TOGETHER_LINES]
    print('表格：共 %d 张，整表同页 %d 张，允许跨页 %d 张（阈值 %d 行）'
          % (len(BB._TABLE_LOG), len(kept), len(span),
             BB.KEEP_TOGETHER_LINES))
    return DOCX_OUT


def render_md():
    """渲染为单一 Markdown 文件（本文件篇幅较短，不另分卷）。"""
    buf = []
    for blk in BLOCKS:
        kind = blk[0]
        if kind in ('cover',):
            buf.append('# 商业计划书设计纲要与撰写样板\n')
            buf.append('> 基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统'
                       '（安步）\n')
            continue
        if kind == 'toc':
            continue
        if kind == 'h1':
            buf.append('\n## ' + blk[1] + '\n')
        elif kind == 'h2':
            buf.append('\n### ' + blk[1] + '\n')
        elif kind == 'h3':
            buf.append('\n#### ' + blk[1] + '\n')
        elif kind == 'p':
            buf.append(blk[1] + '\n')
        elif kind == 'pn':
            buf.append(blk[1] + '\n')
        elif kind == 'b':
            buf.append('- ' + blk[1] + '\n')
        elif kind == 'table':
            _, headers, rows, widths, caption, size, _keep = blk
            if caption:
                buf.append('**' + caption + '**\n')
            buf.append('| ' + ' | '.join(BB._md_cell(h) for h in headers) + ' |')
            buf.append('|' + '|'.join(['---'] * len(headers)) + '|')
            for row in rows:
                buf.append('| ' + ' | '.join(BB._md_cell(c) for c in row) + ' |')
            buf.append('')
        elif kind == 'code':
            buf.append('```')
            buf.append(blk[1])
            buf.append('```\n')
        elif kind == 'center':
            buf.append('> **' + blk[1].replace('\n', ' ') + '**\n')
        elif kind == 'qa':
            _, tag, q, ans = blk
            buf.append('**' + tag + '　' + q + '**\n')
            buf.append('答：' + ans + '\n')
        elif kind == 'pb':
            buf.append('\n---\n')
        elif kind == 'end':
            pass
    with open(MD_OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(buf).rstrip() + '\n')
    return MD_OUT


if __name__ == '__main__':
    print('DOCX  -> ' + render_docx())
    print('MD    -> ' + render_md())
    print('DONE')
