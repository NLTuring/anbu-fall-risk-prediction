# -*- coding: utf-8 -*-
"""
生成《基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统 · 项目简介文档》

设计目标（与用户要求一致）：
  1. 全文档统一使用宋体（含中文、西文、数字、符号），不留任何其他字体；
  2. 本文档是"母版级"说明文档，任何人通读后即可直接撰写挑战杯商业计划书、
     直接搭建答辩 PPT —— 拆解方法见正文 1.4 节与附录 A、附录 B。

本脚本一份内容、双份输出：
  · 基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统-项目简介文档.docx  （主交付物，宋体）
  · docs/*.md                                                  （可编辑的纯文本源）

内容全部写在 tools/content.py 中，本脚本只负责排版与渲染。
"""

import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from content import BLOCKS  # noqa: E402

SONG = '宋体'
ROOT = os.path.dirname(_HERE)
DOCX_OUT = os.path.join(
    ROOT, '基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统-项目简介文档.docx')
DOCS_DIR = os.path.join(ROOT, 'docs')
README_OUT = os.path.join(ROOT, 'README.md')


# ================================================================ DOCX 渲染
_PENDING = {'pb': False, 'gap': False}
# 估算行数不超过该值的表格，整表保持在同一页（避免被拆断后只剩一两行孤悬）
KEEP_TOGETHER_LINES = 21
_TABLE_LOG = []


def _take_gap():
    if _PENDING['gap']:
        _PENDING['gap'] = False
        return 6
    return 0


def set_run_font(run, size=12, bold=False, name=SONG, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), name)


def force_song_everywhere(doc):
    """把所有内置样式与 Normal 的字体统一为宋体，杜绝西文字体回退。"""
    for st in doc.styles:
        try:
            st.font.name = SONG
            rPr = st.element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                rFonts.set(qn(attr), SONG)
        except Exception:
            pass


def add_paragraph(doc, text='', size=12, bold=False, align=None, indent=True,
                  space_before=0, space_after=4, line=1.3, color=None, style=None,
                  keep_next=False):
    para = doc.add_paragraph(style=style)
    fmt = para.paragraph_format
    fmt.line_spacing = line
    fmt.space_before = Pt(space_before + _take_gap())
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_next
    if indent:
        fmt.first_line_indent = Pt(size * 2)
    if align is not None:
        para.alignment = align
    if text:
        add_runs(para, text, size=size, bold=bold, color=color)
    return para


def add_runs(para, text, size=12, bold=False, color=None, italic=False):
    """支持 **...** 行内加粗标记：DOCX 渲染为真实加粗，Markdown 保留原样。"""
    for i, seg in enumerate(str(text).split('**')):
        if not seg:
            continue
        r = para.add_run(seg)
        set_run_font(r, size=size, bold=(bold or i % 2 == 1),
                     color=color, italic=italic)


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    para = doc.add_heading(level=level)
    r = para.add_run(text)
    set_run_font(r, size=sizes.get(level, 12), bold=True, color=RGBColor(0, 0, 0))
    fmt = para.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(14 if level == 1 else (10 if level == 2 else 8))
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Pt(0)
    # 标题与其后的正文或表格保持同页，避免标题孤零零留在页底
    fmt.keep_with_next = True
    if _PENDING['pb']:
        fmt.page_break_before = True
        _PENDING['pb'] = False
    return para


def add_bullet(doc, text, size=12):
    para = doc.add_paragraph(style='List Bullet')
    fmt = para.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(_take_gap())
    fmt.space_after = Pt(3)
    fmt.left_indent = Pt(size * 2)
    fmt.first_line_indent = Pt(0)
    add_runs(para, text, size=size)
    return para


def add_table(doc, headers, rows, widths=None, font_size=9.5, caption=None,
              keep=None):
    if caption:
        cp = doc.add_paragraph()
        cpf = cp.paragraph_format
        cpf.space_before = Pt(6 + _take_gap())
        cpf.space_after = Pt(3)
        cpf.line_spacing = 1.3
        cpf.first_line_indent = Pt(0)
        cpf.keep_with_next = True
        cr = cp.add_run(caption)
        set_run_font(cr, size=10.5, bold=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ''
        hp = hdr[i].paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.line_spacing = 1.0
        hp.paragraph_format.space_after = Pt(2)
        hp.paragraph_format.first_line_indent = Pt(0)
        # 让表头行与其后第一行数据保持同页，避免表头孤零零留在页底
        hp.paragraph_format.keep_with_next = True
        r = hp.add_run(htxt)
        set_run_font(r, size=font_size, bold=True)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:val'), 'clear')
        shade.set(qn('w:fill'), 'EFEFEF')
        hdr[i]._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            cp2 = cells[i].paragraphs[0]
            cp2.paragraph_format.line_spacing = 1.0
            cp2.paragraph_format.space_after = Pt(2)
            cp2.paragraph_format.first_line_indent = Pt(0)
            # 表格单元格同样支持 **行内加粗**（竞品矩阵等需要强调的场合）
            add_runs(cp2, str(val), size=font_size)
    # 表头行跨页重复 + 行不跨页拆分
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    # 小表整表"与下段同页"：避免表格被拆断后末行孤零零留在下一页，
    # 产生只剩一两行数据的近乎空白页。大表仍允许跨页（表头会自动重复）。
    est_lines = _table_lines(headers, rows, widths)
    _TABLE_LOG.append((est_lines, caption or str(headers[0])[:14], len(rows)))
    if keep is True or (keep is None and est_lines <= KEEP_TOGETHER_LINES):
        all_rows = list(t.rows)
        for row in all_rows[:-1]:
            for cell in row.cells:
                for cp3 in cell.paragraphs:
                    cp3.paragraph_format.keep_with_next = True
    # 不在表后插空段落（空段落会被挤到下一页产生整页空白）；
    # 改为记入待处理标记，由下一段通过 space_before 消费这段留白。
    _PENDING['gap'] = True
    return t


def _table_lines(headers, rows, widths):
    """粗估表格占用的文本行数，用于判断能否整表保持在同一页。"""
    total = 0
    for r in [headers] + list(rows):
        row_lines = 1
        for i, cell in enumerate(r):
            w = widths[i] if widths and i < len(widths) else 4.0
            cap = max(4.0, w * 2.6)
            n = 0
            for seg in str(cell).split('\n'):
                n += max(1, int(len(seg) / cap) + (1 if len(seg) % cap else 0))
            row_lines = max(row_lines, n)
        total += row_lines
    return total


def add_code(doc, text, size=10):
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.line_spacing = 1.15
    fmt.space_before = Pt(4 + _take_gap())
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(size * 2)
    # 流程图/架构图是整块字符画，一旦跨页被拆断，
    # 页底会只剩一段无边内容的边框线或一个孤零零的箭头。
    # 因此整块设为"段中不分页"，保证图完整落在一页内。
    fmt.keep_together = True
    r = para.add_run(text)
    set_run_font(r, size=size)
    return para


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Pt(0)
    r = para.add_run()
    set_run_font(r, size=10.5)
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE \\* arabic \\* MERGEFORMAT'
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1)
    r._r.append(instr)
    r._r.append(f2)


def add_toc(doc):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.line_spacing = 1.3
    r = para.add_run()
    set_run_font(r, size=12)
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '【请在 Word 中右键此处，选择"更新域"→"更新整个目录"以生成页码】'
    f3 = OxmlElement('w:fldChar')
    f3.set(qn('w:fldCharType'), 'end')
    for e in (f1, instr, f2, t, f3):
        r._r.append(e)


def add_rule(doc):
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(8)
    fmt.first_line_indent = Pt(0)
    pPr = para._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:color'), '888888')
    pbdr.append(bottom)
    pPr.append(pbdr)


def set_page_start(section, start=1):
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), str(start))
    cols = sectPr.find(qn('w:cols'))
    if cols is not None:
        cols.addprevious(pgNumType)
    else:
        sectPr.append(pgNumType)


def enable_update_fields(doc):
    settings = doc.settings.element
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    settings.append(uf)


def page_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    return sec


def render_docx():
    doc = Document()
    force_song_everywhere(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    for blk in BLOCKS:
        kind = blk[0]

        if kind == 'cover':
            for _ in range(3):
                add_paragraph(doc, '', indent=False, space_after=0)
            add_paragraph(doc, '挑 战 杯 参 赛 项 目', size=14, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=24)
            add_paragraph(doc, '安  步', size=36, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=10)
            add_paragraph(doc, '基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统',
                          size=15, align=WD_ALIGN_PARAGRAPH.CENTER,
                          indent=False, space_after=8)
            add_paragraph(doc, '—— 让跌倒风险在迈出下一步之前被看见 ——',
                          size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                          indent=False, space_after=18)
            add_rule(doc)
            add_paragraph(doc, '本 文 档 用 途', size=12, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=8)
            add_paragraph(
                doc,
                '本文档为项目母版级说明文档，完整收录项目的事实、数据、论证与叙事逻辑。\n'
                '任何人通读本文档后，可直接撰写挑战杯商业计划书，并直接搭建答辩 PPT。\n'
                '具体的拆解方法见正文 1.4 节（本文档使用指南）、'
                '附录 A（商业计划书章节对照表）与附录 B（PPT 页面结构与讲稿要点）。',
                size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                space_after=30, line=1.5)
            for label in ('项目名称：基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统',
                          '建议品牌名：安步（取自“安步当车”，寓意行稳致远）',
                          '申报单位：【待填：学校 / 学院】',
                          '团队名称：【待填：团队名称】',
                          '项目负责人：【待填：姓名】',
                          '指导教师：【待填：姓名 / 职称】',
                          '编制日期：2026 年 10 月'):
                add_paragraph(doc, label, size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                              indent=False, space_after=6)

        elif kind == 'toc':
            page_section(doc)
            add_paragraph(doc, '目  录', size=16, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                          space_after=12)
            add_paragraph(doc, '提示：目录已自动生成。若后续修改正文导致页码变化，'
                               '请右键目录区域选择"更新域"→"更新整个目录"刷新页码。',
                          size=10.5, indent=False, space_after=10,
                          color=RGBColor(0x80, 0x80, 0x80))
            add_toc(doc)

        elif kind == 'h1':
            if not _PENDING['pb'] and not _seen_h1_yet(doc):
                page_section(doc)
                add_page_number_footer(doc.sections[-1])
                set_page_start(doc.sections[-1], 1)
            add_heading(doc, blk[1], level=1)
        elif kind == 'h2':
            add_heading(doc, blk[1], level=2)
        elif kind == 'h3':
            add_heading(doc, blk[1], level=3)
        elif kind == 'p':
            add_paragraph(doc, blk[1], bold=blk[2], indent=blk[3], size=blk[4])
        elif kind == 'pn':
            add_paragraph(doc, blk[1], bold=blk[2], indent=False, size=blk[3])
        elif kind == 'b':
            add_bullet(doc, blk[1])
        elif kind == 'table':
            _, headers, rows, widths, caption, size, keep = blk
            add_table(doc, headers, rows, widths, size, caption, keep)
        elif kind == 'code':
            add_code(doc, blk[1], size=blk[2])
        elif kind == 'center':
            add_paragraph(doc, blk[1], size=blk[2], bold=blk[3],
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                          space_before=6, space_after=8)
        elif kind == 'qa':
            _, tag, q, ans = blk
            qp = doc.add_paragraph()
            qf = qp.paragraph_format
            qf.line_spacing = 1.3
            qf.space_before = Pt(8)
            qf.space_after = Pt(3)
            qf.first_line_indent = Pt(0)
            # 问题与其答案保持同页，避免问题孤行在页底而答案翻页
            qf.keep_with_next = True
            add_runs(qp, tag + '　' + q, size=12, bold=True)
            ap = doc.add_paragraph()
            af = ap.paragraph_format
            af.line_spacing = 1.3
            af.space_after = Pt(4)
            af.left_indent = Pt(24)
            af.first_line_indent = Pt(0)
            add_runs(ap, '答：' + ans, size=11.5)
        elif kind == 'pb':
            _PENDING['pb'] = True
        elif kind == 'mdnew':
            pass
        elif kind == 'end':
            add_rule(doc)
            add_paragraph(doc, '（文档结束）', size=11,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
                          space_before=10)

    enable_update_fields(doc)
    doc.save(DOCX_OUT)
    kept = [c for n, c, r in _TABLE_LOG if n <= KEEP_TOGETHER_LINES]
    span = [c for n, c, r in _TABLE_LOG if n > KEEP_TOGETHER_LINES]
    print('表格：共 %d 张，整表同页 %d 张，允许跨页 %d 张（阈值 %d 行）'
          % (len(_TABLE_LOG), len(kept), len(span), KEEP_TOGETHER_LINES))
    with open(os.path.join(ROOT, 'tools', '_tables.txt'), 'w',
              encoding='utf-8') as f:
        for n, c, r in sorted(_TABLE_LOG, reverse=True):
            f.write('%3d 行  %2d 行数据  %s\n' % (n, r, c))
    return DOCX_OUT


_H1_SEEN = {'v': False}


def _seen_h1_yet(doc):
    """第一章之前需要新建正文节（目录单独一节，正文页码从 1 开始）。"""
    if _H1_SEEN['v']:
        return True
    _H1_SEEN['v'] = True
    return False


# ================================================================ Markdown 渲染
def render_md():
    """把同一份内容渲染为可编辑的 Markdown 分卷文件，便于后续修改。"""
    if not os.path.isdir(DOCS_DIR):
        os.makedirs(DOCS_DIR)
    written = []
    current_path = None
    current_title = None
    buf = []

    def flush():
        if current_path is None:
            return
        os.makedirs(os.path.dirname(current_path), exist_ok=True)
        head = '# ' + current_title + '\n\n'
        with open(current_path, 'w', encoding='utf-8') as f:
            f.write(head + '\n'.join(buf).rstrip() + '\n')
        written.append(current_path)

    for blk in BLOCKS:
        kind = blk[0]
        if kind == 'cover' or kind == 'toc':
            continue
        if kind == 'mdnew':
            flush()
            current_path = os.path.join(ROOT, blk[1])
            current_title = blk[2]
            buf = []
            continue
        if current_path is None:
            continue

        if kind == 'h1':
            buf.append('\n## ' + blk[1] + '\n')
        elif kind == 'h2':
            buf.append('\n### ' + blk[1] + '\n')
        elif kind == 'h3':
            buf.append('\n#### ' + blk[1] + '\n')
        elif kind == 'p':
            buf.append(blk[1] + '\n')
        elif kind == 'pn':
            buf.append(blk[1] + '\n')
        elif kind == 'b':
            buf.append('- ' + blk[1] + '\n')
        elif kind == 'table':
            _, headers, rows, widths, caption, size, _keep = blk
            if caption:
                buf.append('**' + caption + '**\n')
            buf.append('| ' + ' | '.join(_md_cell(h) for h in headers) + ' |')
            buf.append('|' + '|'.join(['---'] * len(headers)) + '|')
            for row in rows:
                buf.append('| ' + ' | '.join(_md_cell(c) for c in row) + ' |')
            buf.append('')
        elif kind == 'code':
            buf.append('```')
            buf.append(blk[1])
            buf.append('```\n')
        elif kind == 'center':
            buf.append('> **' + blk[1].replace('\n', ' ') + '**\n')
        elif kind == 'qa':
            _, tag, q, ans = blk
            buf.append('**' + tag + '　' + q + '**\n')
            buf.append('答：' + ans + '\n')
        elif kind == 'pb':
            buf.append('\n---\n')
        elif kind == 'end':
            pass

    flush()
    return written


def _md_cell(v):
    return str(v).replace('\n', '<br>').replace('|', '\\|')


# ================================================================ README
def render_readme():
    lines = [
        '# 安步 · 基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统',
        '',
        '> 一句话：让跌倒风险**在迈出下一步之前**被看见。',
        '',
        '本仓库为挑战杯参赛项目的材料母版。核心交付物是**项目简介文档**——',
        '一份读完即可直接撰写商业计划书、直接搭建答辩 PPT 的说明文档。',
        '',
        '## 交付物',
        '',
        '| 文件 | 说明 |',
        '|---|---|',
        '| [基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统-项目简介文档.docx]'
        '(基于改进 YOLO 与步态分析的老龄人群跌倒风险预测系统-项目简介文档.docx) '
        '| **主交付物**，全文档宋体，含封面 / 自动目录 / 十三章正文 / 五个附录 |',
        '| [商业计划书设计纲要与样板.docx](商业计划书设计纲要与样板.docx) '
        '| **配套交付物**，宋体，含商业计划书 15 章设计纲要 + '
        '可直接在 Word 里扩写的简略样板 |',
        '| [商业计划书设计纲要与样板.md](商业计划书设计纲要与样板.md) '
        '| 上述文档的可编辑纯文本源 |',
        '| [docs/01-项目速览与痛点.md](docs/01-项目速览与痛点.md) '
        '| 项目速览、五种长度简介、核心数据卡、'
        '本文档使用指南、三浪叠加与三层痛点 |',
        '| [docs/02-产品与技术方案.md](docs/02-产品与技术方案.md) '
        '| 产品定位与形态、硬件配置、功能清单、产品边界、'
        '四级映射与四个缺口、四个核心构造、五层架构、技术边界 |',
        '| [docs/03-创新点与竞品分析.md](docs/03-创新点与竞品分析.md) '
        '| 四条创新点的四段式论证、三类竞品全景、竞品矩阵、'
        '差异化定位与四层壁垒 |',
        '| [docs/04-市场商业模式与团队规划.md](docs/04-市场商业模式与团队规划.md) '
        '| 市场分层与自下而上测算、商业模式画布、四层收入、'
        '单位经济、财务、团队、路线图、风险 |',
        '| [docs/05-成果与BP·PPT拆解指南.md](docs/05-成果与BP·PPT拆解指南.md) '
        '| **附录 A–E**：商业计划书章节对照表、PPT 页面结构与讲稿、'
        '答辩问答库、素材证据清单、数据核实要求 |',
        '| [tools/content.py](tools/content.py) '
        '| **项目简介文档的全部内容**，用一套小型 DSL 描述，'
        '改内容只改这一个文件 |',
        '| [tools/bp_content.py](tools/bp_content.py) '
        '| **商业计划书纲要与样板的全部内容**，同样用 DSL 描述 |',
        '| [tools/build_brief.py](tools/build_brief.py) · '
        '[tools/build_bp.py](tools/build_bp.py) '
        '| 生成脚本：一份内容同时输出 docx 与 Markdown，改内容后重跑即可 |',
        '| [tools/verify_brief.py](tools/verify_brief.py) '
        '| 字体与结构自检：逐 run 校验宋体、残留标记、目录域与标题层级；'
        '可传 docx 文件名以校验任一份文档 |',
        '| [tools/check_quotes.py](tools/check_quotes.py) '
        '| 语法预检：找出内容里引号不成对的行（中文内容最易踩的坑） |',
        '| [tools/export_pdf.py](tools/export_pdf.py) · '
        '[tools/render_png.py](tools/render_png.py) '
        '| 用 Word COM 导出 PDF 并逐页渲染 PNG，用于版式验收 |',
        '| [preview/brief.pdf](preview/brief.pdf) · `preview/pages/` '
        '| 项目简介文档的 PDF 与逐页图片（72 页） |',
        '| [preview/bp.pdf](preview/bp.pdf) · `preview/bp_pages/` '
        '| 商业计划书纲要与样板的 PDF 与逐页图片（23 页正文） |',
        '',
        '## 怎么用',
        '',
        '1. **写商业计划书**：打开《商业计划书设计纲要与样板.docx》。'
        '上半部分是逐章写作纲要（每章要立住哪一句结论、写多长、'
        '用什么图、素材从哪来、评委在哪扣分），'
        '下半部分是可直接扩写的 15 章简略样板——'
        '把【待填】【待实测】补上自己的数据即可成稿。',
        '2. **做答辩 PPT**：打开项目简介文档的附录 B，按页面结构表逐页制作，'
        '表中已给出标题、核心内容、视觉素材、讲稿要点与时长。',
        '3. **准备答辩问答**：项目简介文档附录 C 的 Q1–Q16 逐条演练，'
        '其中 Q10 必须背诵。',
        '4. **推进项目**：项目简介文档附录 D 是核验表，按优先级自上而下完成'
        '（★ 最高优先级的五项是成本最低、取信效果最好的证据）。',
        '5. **提交前自查**：项目简介文档附录 E 列出所有待核实数据，'
        '必须逐项清零。',
        '',
        '## 项目要点速记',
        '',
        '- **做什么**：用 RK3576 边缘盒 + 普通 USB 摄像头，'
        '在养老机构走廊与家庭客厅里完成无穿戴步态量化与跌倒风险预测。',
        '- **为什么是现在**：需求侧 3.23 亿老年人口、'
        '供给侧 767.9 万张床位对 72.2 万从业人员、'
        '技术侧国产 NPU 与视觉模型同时平价化。',
        '- **四条创新**：① 面向步态测量的姿态估计改进；'
        '② 单目弱标定的物理尺度步态量化；'
        '③ 以步态变异性为核心的连续趋势风险建模；'
        '④ 视频不出设备的端侧隐私安全架构。',
        '- **最大的差异化**：别人测“有没有摔倒”，'
        '我们测“会不会摔倒”；别人把视频传到云上，我们让视频不出设备。',
        '',
        '## 修改方式',
        '',
        '内容全部位于 `tools/content.py`，用一套小型 DSL 描述',
        '（`h1/h2/p/b/tbl/code` 等），排版逻辑位于 `tools/build_brief.py`。改完重跑：',
        '',
        '```bash',
        'python tools/build_brief.py',
        '```',
        '',
        '脚本会同时刷新 docx 与 docs/ 下的 Markdown 分卷，两者内容天然一致。',
        '全文档字体统一为**宋体**（含西文与数字），由 `force_song_everywhere()`',
        '与逐 run 的 `rFonts` 设置双重保证。',
        '',
        '排版上有两条自动化规则，改内容后无需手工调版：',
        '',
        '- **小表格整表同页**：估算行数 ≤ 21 的表格自动保持在同一页，',
        '  避免被拆断后只剩一两行孤悬在下一页；需要强制或禁止可对 `tbl()`',
        '  传 `keep=True/False`。',
        '- **标题与表格不孤行**：标题、表格表头行、附录 C 的问题段落均设为',
        '  “与下段同页”，避免标题或问题落在页底而内容翻页。',
        '',
        '改完内容后建议依次执行：',
        '',
        '```bash',
        'python tools/check_quotes.py tools/content.py   # 引号预检',
        'python tools/build_brief.py                     # 生成 docx + md',
        'python tools/verify_brief.py                    # 字体结构自检',
        'python tools/export_pdf.py && python tools/render_png.py   # 导出并出图',
        '```',
        '',
        '说明：本文档每一章都从新的一页开始，因此每章**最后一页**留白偏多，',
        '这是该排版约定的正常结果（正式申报材料普遍如此），并非缺陷。',
        '若希望消除留白，把对应章节的 `pb()` 标记删去即可让正文连续排版。',
        '',
        '## 待办（优先级从高到低）',
        '',
        '- [ ] ★ 抓包验证：证明系统运行期间对外流量不含视频数据（含录屏）',
        '- [ ] ★ 消融实验：量化四项姿态估计改进各自的贡献',
        '- [ ] ★ 重复性验证：测量系统自身引入的变异，'
        '证明“测到的变异来自老人而非算法”',
        '- [ ] ★ 端侧样机跑通全流程，并录制离线兜底演示视频',
        '- [ ] ★ 落实临床合作单位（医院康复科或老年科），'
        '启动金标准一致性验证',
        '- [ ] 实测样机 BOM 成本；向至少 3 家单位询价临床步态设备',
        '- [ ] 完成 3–5 家养老机构访谈，采集纠纷成本与巡房频次等 ROI 原始数据',
        '- [ ] 申请伦理审查；准备知情同意书与隐私设计对照表',
        '- [ ] 完成附录 E 中全部标注 ★ 的数据核实',
        '- [ ] 确认品牌名（现用“安步”为建议名）与参赛赛道',
        '',
    ]
    with open(README_OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    return README_OUT



# ================================================================ 主流程
if __name__ == '__main__':
    out_docx = render_docx()
    out_mds = render_md()
    out_readme = render_readme()
    print('DOCX  -> ' + out_docx)
    for m in out_mds:
        print('MD    -> ' + m)
    print('README-> ' + out_readme)
    print('DONE')
